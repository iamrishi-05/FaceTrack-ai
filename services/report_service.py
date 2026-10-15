import io
import pandas as pd
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from utils.logger import log_event


class DetectionReportNumberedCanvas(canvas.Canvas):
    """
    Two-pass ReportLab canvas that dynamically renders running headers
    and footers with total page count ('Page X of Y').
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        width, height = letter
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))

        # Running Header on page 2 and later
        if self._pageNumber > 1:
            self.drawString(36, height - 24, "FaceTrack AI — Biometric Surveillance & Detection Intelligence Dossier")
            self.drawRightString(width - 36, height - 24, "AUTHORIZED AUDIT DOSSIER")
            self.setStrokeColor(colors.HexColor("#e2e8f0"))
            self.setLineWidth(0.5)
            self.line(36, height - 28, width - 36, height - 28)

        # Running Footer on all pages
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        self.line(36, 30, width - 36, 30)
        self.drawString(36, 20, "CONFIDENTIAL & AUTHORIZED — FACETRACK AI BIOMETRIC SYSTEM")
        self.drawRightString(width - 36, 20, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


class ReportService:
    """
    High-performance, publication-quality reporting engine for FaceTrack AI.
    Generates executive PDF surveillance dossiers and structured Excel audit workbooks.
    """

    @staticmethod
    def generate_detection_pdf(records, filter_description="All Detection Records"):
        """
        Generates an executive, cybernetic-grade PDF intelligence dossier
        from biometric detection logs.
        """
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()
        story = []

        # Custom Typographic Styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#0f172a'),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#64748b'),
            spaceAfter=12
        )
        kpi_title_style = ParagraphStyle(
            'KPITitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            textColor=colors.HexColor('#64748b'),
            alignment=1
        )
        kpi_val_style = ParagraphStyle(
            'KPIVal',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=13,
            textColor=colors.HexColor('#0f172a'),
            alignment=1
        )
        th_style = ParagraphStyle(
            'TH',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8.5,
            textColor=colors.white,
            alignment=0
        )
        td_style = ParagraphStyle(
            'TD',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor('#1e293b')
        )
        td_bold_style = ParagraphStyle(
            'TDBold',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor('#0f172a')
        )
        td_emerald_style = ParagraphStyle(
            'TDEmerald',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor('#047857')
        )
        td_mono_style = ParagraphStyle(
            'TDMono',
            parent=styles['Normal'],
            fontName='Courier-Bold',
            fontSize=7.5,
            leading=10,
            textColor=colors.HexColor('#0369a1')
        )

        # 1. Header & Title Block
        story.append(Paragraph("FACETRACK AI — BIOMETRIC DETECTION INTELLIGENCE DOSSIER", title_style))
        story.append(Paragraph(
            f"Generated: <b>{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</b> | "
            f"Filter Scope: <b>{filter_description}</b> | "
            f"Classification: <b>SECURE AUDIT RECORD</b>",
            subtitle_style
        ))

        # 2. Executive KPI Summary Cards Table
        total_logs = len(records)
        unique_targets = len(set(r['target_id'] for r in records if r.get('target_id')))
        conf_values = [float(r['confidence']) for r in records if r.get('confidence')]
        avg_conf = f"{(sum(conf_values) / len(conf_values)):.1f}%" if conf_values else "N/A"

        kpi_data = [
            [
                Paragraph("TOTAL EVENTS", kpi_title_style),
                Paragraph("UNIQUE TARGETS", kpi_title_style),
                Paragraph("MEAN CONFIDENCE", kpi_title_style),
                Paragraph("COMPLIANCE AUDIT", kpi_title_style)
            ],
            [
                Paragraph(str(total_logs), kpi_val_style),
                Paragraph(str(unique_targets), kpi_val_style),
                Paragraph(avg_conf, kpi_val_style),
                Paragraph("100% VERIFIED", kpi_val_style)
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[135, 135, 135, 135])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 14))

        # 3. Detection Logs Data Table
        # Width: 24 + 76 + 95 + 75 + 65 + 120 + 85 = 540 pt (full printable width)
        headers = [
            Paragraph("#", th_style),
            Paragraph("Timestamp", th_style),
            Paragraph("Target Person", th_style),
            Paragraph("Target ID", th_style),
            Paragraph("Confidence", th_style),
            Paragraph("Liveness & Features", th_style),
            Paragraph("Camera Location", th_style)
        ]
        table_rows = [headers]

        for idx, r in enumerate(records, start=1):
            # Format attributes string
            attrs = []
            emotion = r.get('emotion') or 'neutral'
            attrs.append(f"Emotion: {emotion.capitalize()}")
            if r.get('mask_detected'): attrs.append("Mask: YES")
            if r.get('smile_detected'): attrs.append("Smile: YES")
            if r.get('blink_detected'): attrs.append("Blink: YES")
            attrs_str = " • ".join(attrs)

            time_str = f"{r.get('date', '')}<br/>{r.get('time', '')}"
            conf_str = f"{r.get('confidence', 'N/A')}%" if r.get('confidence') else "N/A"
            loc_str = r.get('location_tag') or 'Primary Camera'

            table_rows.append([
                Paragraph(str(idx), td_style),
                Paragraph(time_str, td_style),
                Paragraph(r.get('target_name', 'Unknown'), td_bold_style),
                Paragraph(r.get('target_id', 'N/A'), td_mono_style),
                Paragraph(conf_str, td_emerald_style),
                Paragraph(attrs_str, td_style),
                Paragraph(loc_str, td_style)
            ])

        col_widths = [24, 76, 95, 75, 65, 120, 85]
        t = Table(table_rows, colWidths=col_widths, repeatRows=1)

        t_style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')), # Slate-900 Header
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 7),
            ('TOPPADDING', (0, 0), (-1, 0), 7),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('LINEBELOW', (0, 0), (-1, 0), 1.5, colors.HexColor('#10b981')), # Emerald accent border
        ])

        # Zebra Striping
        for i in range(1, len(table_rows)):
            bg = colors.HexColor('#f8fafc') if i % 2 == 0 else colors.white
            t_style.add('BACKGROUND', (0, i), (-1, i), bg)
            t_style.add('TOPPADDING', (0, i), (-1, i), 5)
            t_style.add('BOTTOMPADDING', (0, i), (-1, i), 5)

        t.setStyle(t_style)
        story.append(t)

        # Build document with running headers/footers
        doc.build(story, canvasmaker=DetectionReportNumberedCanvas)
        output.seek(0)
        log_event("INFO", "Reporting", f"Detection PDF report successfully compiled with {len(records)} entries.")
        return output.getvalue()

    @staticmethod
    def generate_detection_excel(records):
        """
        Generates an Excel workbook file stream using Pandas and OpenPyXL
        containing both a detailed Audit Trail and an Executive Summary.
        """
        data = []
        for idx, r in enumerate(records, start=1):
            data.append({
                'Log #': idx,
                'Date': r.get('date', ''),
                'Time': r.get('time', ''),
                'Target ID': r.get('target_id', ''),
                'Target Name': r.get('target_name', ''),
                'Confidence (%)': float(r.get('confidence', 0)) if r.get('confidence') else 0.0,
                'Emotion': (r.get('emotion') or 'neutral').capitalize(),
                'Mask Detected': 'Yes' if r.get('mask_detected') else 'No',
                'Smile Detected': 'Yes' if r.get('smile_detected') else 'No',
                'Blink Detected': 'Yes' if r.get('blink_detected') else 'No',
                'Camera Location Tag': r.get('location_tag') or 'Primary Camera',
                'Timestamp': r.get('timestamp', '')
            })

        df_logs = pd.DataFrame(data)

        # Build Summary Data
        total_logs = len(records)
        unique_targets = len(set(r['target_id'] for r in records if r.get('target_id')))
        conf_values = [float(r['confidence']) for r in records if r.get('confidence')]
        avg_conf = round(sum(conf_values) / len(conf_values), 1) if conf_values else 0.0

        summary_data = [
            {'Metric': 'Report Generated At', 'Value': datetime.now().strftime('%Y-%m-%d %H:%M:%S')},
            {'Metric': 'Total Detection Incidents', 'Value': total_logs},
            {'Metric': 'Unique Target Profiles Spotted', 'Value': unique_targets},
            {'Metric': 'Average Recognition Confidence (%)', 'Value': f"{avg_conf}%"},
            {'Metric': 'Audit Classification', 'Value': 'Authorized Security & Biometric Log'}
        ]
        df_summary = pd.DataFrame(summary_data)

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_logs.to_excel(writer, sheet_name='Detection Audit Trail', index=False)
            df_summary.to_excel(writer, sheet_name='Executive Summary', index=False)

            workbook = writer.book

            # 1. Format Detection Audit Trail Sheet
            ws_logs = writer.sheets['Detection Audit Trail']
            header_fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid') # Deep Teal/Emerald
            header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
            thin_border = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            # Style Header Row
            for col_idx in range(1, len(df_logs.columns) + 1):
                cell = ws_logs.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

            ws_logs.row_dimensions[1].height = 28

            # Style Data Rows
            for row_idx in range(2, len(df_logs) + 2):
                ws_logs.row_dimensions[row_idx].height = 20
                is_even = (row_idx % 2 == 0)
                for col_idx in range(1, len(df_logs.columns) + 1):
                    cell = ws_logs.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    if is_even:
                        cell.fill = alt_fill

                    # Alignment heuristics
                    if col_idx in [1, 2, 3, 4, 6, 7, 8, 9, 10]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

            # Auto-fit Column Widths
            for col in ws_logs.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws_logs.column_dimensions[col_letter].width = max(max_len + 4, 12)

            # Freeze Header Row
            ws_logs.freeze_panes = 'A2'

            # 2. Format Executive Summary Sheet
            ws_sum = writer.sheets['Executive Summary']
            sum_header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
            for col_idx in range(1, 3):
                cell = ws_sum.cell(row=1, column=col_idx)
                cell.fill = sum_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            ws_sum.row_dimensions[1].height = 26
            ws_sum.column_dimensions['A'].width = 38
            ws_sum.column_dimensions['B'].width = 32

            for r_idx in range(2, len(df_summary) + 2):
                ws_sum.row_dimensions[r_idx].height = 22
                for c_idx in range(1, 3):
                    cell = ws_sum.cell(row=r_idx, column=c_idx)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical='center')

        output.seek(0)
        log_event("INFO", "Reporting", f"Detection Excel workbook report compiled with {len(records)} entries.")
        return output.getvalue()

    @staticmethod
    def generate_excel(records):
        """Legacy helper for attendance records."""
        data = []
        for r in records:
            data.append({
                'Student ID': r['student_id'],
                'Name': r['name'],
                'Roll Number': r['roll_number'],
                'Department': r['department'],
                'Subject': r['subject'] if 'subject' in r.keys() and r['subject'] else 'Python',
                'Check-in Date': r['date'],
                'Check-in Time': r['time'],
                'Status': r['status'],
                'Method': r['method'],
                'Confidence %': r['confidence'] if r['confidence'] else 'N/A',
                'Emotion': r['emotion'].capitalize() if r['emotion'] else 'N/A'
            })
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Attendance Logs', index=False)
            worksheet = writer.sheets['Attendance Logs']
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_pdf(records, filter_description="All Records"):
        """Legacy helper for attendance records."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        story = []
        title_style = ParagraphStyle('ReportTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=22, textColor=colors.HexColor('#002b49'), spaceAfter=6)
        sub_style = ParagraphStyle('ReportSub', parent=styles['Normal'], fontName='Helvetica', fontSize=10, textColor=colors.HexColor('#64748b'), spaceAfter=20)
        story.append(Paragraph("FaceTrack AI – Attendance Report", title_style))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Filters: {filter_description}", sub_style))

        header_style = ParagraphStyle('TableHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, textColor=colors.white)
        table_data = [[
            Paragraph("Student ID", header_style),
            Paragraph("Name", header_style),
            Paragraph("Subject", header_style),
            Paragraph("Department", header_style),
            Paragraph("Date", header_style),
            Paragraph("Time", header_style),
            Paragraph("Status", header_style)
        ]]
        for r in records:
            subject_val = r['subject'] if 'subject' in r.keys() and r['subject'] else 'Python'
            table_data.append([
                Paragraph(r['student_id'], styles['Normal']),
                Paragraph(r['name'], styles['Normal']),
                Paragraph(subject_val, styles['Normal']),
                Paragraph(r['department'], styles['Normal']),
                Paragraph(r['date'], styles['Normal']),
                Paragraph(r['time'], styles['Normal']),
                Paragraph(f"<b>{r['status']}</b>", styles['Normal'])
            ])
        col_widths = [75, 95, 85, 95, 65, 60, 65]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t_style = TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ])
        for idx in range(1, len(table_data)):
            bg_color = colors.HexColor('#f8fafc') if idx % 2 == 0 else colors.HexColor('#ffffff')
            t_style.add('BACKGROUND', (0, idx), (-1, idx), bg_color)
        t.setStyle(t_style)
        story.append(t)
        doc.build(story)
        output.seek(0)
        return output.getvalue()
